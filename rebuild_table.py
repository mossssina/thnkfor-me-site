"""
Полный ребилд niches_analysis.xlsx:
- сортировка по Приоритет → Итог (убывание)
- единое чередующееся оформление строк
- новый столбец 14: Предлагаемые услуги
- AutoFilter на строку заголовков
- закреплённые строки 1–2
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILE = "niches_analysis.xlsx"

# ── Цвета категорий (как в оригинале) ──────────────────────────────────────
CAT_COLORS = {
    "HoReCa":          "BF360C",
    "Beauty":          "880E4F",
    "Beauty/Медицина": "6A1B9A",
    "Медицина":        "1565C0",
    "Фитнес":          "1B5E20",
    "Образование":     "E65100",
    "Ветеринария":     "33691E",
    "B2B":             "4A148C",
    "Гостеприимство":  "6D4C41",
    "Event и творчество": "AD1457",
    "Авто":            "212121",
    "Дом и быт":       "455A64",
    "Недвижимость":    "01579B",
    "Коучинг":         "37474F",
    "Туризм":          "00695C",
    "E-commerce":      "E65100",
}

# ── Цвета по значению ───────────────────────────────────────────────────────
def fill(hex_color):
    if not hex_color or hex_color == "000000":
        return PatternFill(fill_type=None)
    return PatternFill(fill_type="solid", fgColor=hex_color)

def stars_color(s):
    n = s.count("★") if s else 0
    if n >= 5: return "1B5E20"
    if n == 4: return "2E7D32"
    if n == 3: return "E65100"
    return "B71C1C"

def pay_color(v):
    return {"Высокая": "E8F5E9", "Средняя": "FFF3E0", "Низкая": "FFEBEE"}.get(v, "FFFFFF")

def freq_color(v):
    return {"Высокая": "E8F5E9", "Средняя": "FFF8E1", "Сезонная": "FFFDE7", "Низкая": "FFFDE7"}.get(v, "FFFFFF")

def comp_color(v):
    return {"Высокая": "FFEBEE", "Средняя": "FFF8E1", "Низкая": "E8F5E9"}.get(v, "FFFFFF")

def score_color(v):
    if v is None: return "FFFFFF"
    try:
        n = int(v)
        if n >= 9: return "1B5E20"
        if n >= 8: return "2E7D32"
        if n >= 7: return "2E7D32"
        if n >= 6: return "E65100"
        return "B71C1C"
    except: return "FFFFFF"

def prio_color(v):
    return {"P0": "1B5E20", "P1": "E65100", "P2": "37474F"}.get(str(v), "FFFFFF")

def scenario_color(v):
    if v and "А" in str(v) and "Б" in str(v): return "E8EAF6"
    if v == "А": return "E3F2FD"
    if v == "Б": return "E8F5E9"
    return "F3E5F5"

# ── Предлагаемые услуги ─────────────────────────────────────────────────────
SERVICES = {
    # HoReCa
    "Рестораны":                    "AI-контент, автопостинг, бот-бронирование, интеграция iiko",
    "Кафе":                         "AI-контент, автопостинг, TG-бот, Яндекс.Карты",
    "Бары и пабы":                  "AI-контент, автопостинг, TG-бот",
    "Кофейни":                      "AI-контент, автопостинг, TG-бот",
    "Доставка еды (свой бренд)":    "AI-контент, автопостинг, бот-заказ, интеграция iiko",
    "Веганские и здоровые кафе":    "AI-контент, автопостинг, TG-бот",
    "Пекарни и кондитерские":       "AI-контент, автопостинг, TG-бот",
    "Кейтеринг":                    "AI-контент, TG-бот, автоматизация заявок",
    "Фуд-траки":                    "AI-контент, TG-бот",
    "Тёмные кухни (dark kitchen)":  "AI-контент, автопостинг, бот-заказ",
    # Beauty
    "Салоны красоты":               "AI-контент, автопостинг, бот-запись, CRM-интеграция",
    "Барбершопы":                   "AI-контент, автопостинг, бот-запись",
    "Ногтевые студии":              "AI-контент, автопостинг, бот-запись",
    "SPA-салоны":                   "AI-контент, автопостинг, бот-запись, CRM-интеграция",
    "Студии лазерной эпиляции":     "AI-контент, автопостинг, бот-запись",
    "Перманентный макияж":          "AI-контент, TG-бот, автопостинг",
    "Бровисты / лешмейкеры":       "AI-контент, TG-бот, бот-запись",
    "Массажные студии":             "AI-контент, автопостинг, бот-запись",
    "Татуировки и пирсинг":         "AI-контент, TG-бот",
    "Студии загара":                "AI-контент, TG-бот",
    "Груминг-салоны":               "AI-контент, автопостинг, бот-запись",
    "Косметологические клиники":    "AI-контент, автопостинг, бот-запись, CRM-интеграция",
    # Медицина
    "Стоматологические клиники":    "AI-контент, бот-запись, CRM-интеграция, Prodoctorov",
    "Частные клиники (много.)":     "AI-контент, бот-запись, CRM-интеграция, MedMe",
    "Дерматологические клиники":    "AI-контент, бот-запись, CRM-интеграция",
    "Пластическая хирургия":        "AI-контент, бот-запись, CRM-интеграция",
    "Офтальмологические клиники":   "AI-контент, бот-запись, CRM-интеграция",
    "Гинекологические клиники":     "AI-контент, бот-запись, CRM-интеграция",
    "Реабилитационные центры":      "AI-контент, бот-запись, CRM-интеграция",
    "Медицинские лаборатории":      "AI-контент, TG-бот, CRM-интеграция",
    "Ветеринарные клиники":         "AI-контент, бот-запись, CRM-интеграция",
    # Фитнес
    "Фитнес-клубы":                 "AI-контент, автопостинг, бот-запись, CRM-интеграция",
    "Йога-студии":                  "AI-контент, автопостинг, бот-запись",
    "Пилатес-студии":               "AI-контент, автопостинг, бот-запись",
    "Кроссфит-боксы":               "AI-контент, автопостинг, бот-запись",
    "Студии растяжки":              "AI-контент, автопостинг, бот-запись",
    "Танцевальные студии":          "AI-контент, автопостинг, бот-запись",
    "Единоборства (бокс, MMA)":    "AI-контент, автопостинг, бот-запись",
    "Батутные центры":              "AI-контент, автопостинг, TG-бот",
    "Скалодромы":                   "AI-контент, TG-бот",
    "Конные клубы":                 "AI-контент, TG-бот",
    "Теннисные клубы":              "AI-контент, TG-бот, бот-запись",
    "Бассейны (частные)":           "AI-контент, TG-бот, бот-запись",
    # Образование
    "Детские центры развития":      "AI-контент, автопостинг, TG-бот, автоворонка",
    "Языковые школы":               "AI-контент, автопостинг, TG-бот, автоворонка",
    "Музыкальные школы":            "AI-контент, TG-бот",
    "Художественные школы":         "AI-контент, TG-бот",
    "Робототехника / STEM":         "AI-контент, TG-бот, автоворонка",
    "Частные детские сады":         "AI-контент, TG-бот",
    "Онлайн-курсы / EdTech":        "AI-контент, автопостинг, TG-бот, автоворонка",
    "Частные школы":                "AI-контент, TG-бот",
    "Шахматные клубы":              "AI-контент, TG-бот",
    "Репетиторы (частные)":         "AI-контент, TG-бот",
    # Гостеприимство
    "Отели (бутик, 3–4*)":         "AI-контент, автопостинг, TG-бот, CRM-интеграция",
    "Санатории и курорты":          "AI-контент, TG-бот, автоматизация бронирования",
    "Глэмпинг":                     "AI-контент, автопостинг, TG-бот",
    "Загородные базы отдыха":       "AI-контент, TG-бот",
    "Хостелы":                      "AI-контент, TG-бот",
    "Посуточная аренда (AirBnB)":  "AI-контент, TG-бот, автоматизация бронирования",
    # E-commerce существующие
    "Инт. магазины (спортивные)":   "AI-контент, автопостинг, TG-бот",
    "Инт. магазины (мода/одежда)":  "AI-контент, автопостинг, TG-бот",
    "Инт. магазины (косметика)":    "AI-контент, автопостинг, TG-бот",
    "Инт. магазины (детские тов.)": "AI-контент, автопостинг, TG-бот",
    "Зоотовары онлайн":             "AI-контент, автопостинг, TG-бот",
    "Зоомагазины":                  "AI-контент, автопостинг, TG-бот",
    # E-commerce новые (парсинг)
    "Маркетплейс-продавцы (WB / Ozon)":    "Парсинг каталога WB/Ozon, AI-контент, автопостинг, TG-бот",
    "Инт. магазины (мебель и декор)":       "Парсинг каталога, AI-контент, автопостинг, TG-бот",
    "Инт. магазины (ювелирные украшения)":  "Парсинг каталога, AI-контент, автопостинг, TG-бот",
    "Инт. магазины (электроника и гаджеты)":"Парсинг каталога, AI-контент, автопостинг, TG-бот",
    "Инт. магазины (товары для дома)":      "Парсинг каталога, AI-контент, автопостинг, TG-бот",
    "Фермерские магазины / натуральная еда":"Парсинг каталога, AI-контент, автопостинг, TG-бот",
    "Аптеки и товары для здоровья онлайн":  "Парсинг каталога, AI-контент, автопостинг, TG-бот",
    "Книжные магазины онлайн":              "Парсинг каталога, AI-контент, автопостинг, TG-бот",
    # Авто
    "Автосервисы":                  "AI-контент, TG-бот, автоматизация записи",
    "Аренда автомобилей":           "AI-контент, TG-бот, автоматизация бронирования",
    "Автосалоны (независимые)":     "AI-контент, TG-бот, CRM-интеграция",
    "Детейлинг-студии":             "AI-контент, TG-бот, бот-запись",
    "Автомойки":                    "AI-контент, TG-бот",
    "Шинные центры":                "AI-контент, TG-бот",
    # Дом и быт
    "Клининговые компании":         "AI-контент, TG-бот, автоматизация заявок",
    "Дизайнеры интерьера":          "AI-контент, TG-бот, автоматизация лидов",
    "Ремонт квартир (бригады)":     "AI-контент, TG-бот, автоматизация заявок",
    "Ландшафтные дизайнеры":        "AI-контент, TG-бот",
    "Установка окон и дверей":      "AI-контент, TG-бот",
    "Натяжные потолки":             "AI-контент, TG-бот",
    # Недвижимость
    "Коворкинги":                   "AI-контент, TG-бот, автоматизация бронирования",
    "Застройщики (малые)":          "AI-контент, TG-бот, CRM-интеграция",
    "Агентства недвижимости":       "AI-контент, TG-бот, CRM-интеграция",
    "Частные риелторы":             "AI-контент, TG-бот",
    # B2B
    "SMM-агентства":                "AI-контент, TG-бот (команда), автопостинг, парсинг",
    "Маркетинговые агентства":      "AI-контент, TG-бот (команда), автопостинг",
    "Онлайн-школы (B2B-контент)":  "AI-контент, автопостинг, TG-бот, автоворонка",
    "Юридические фирмы":            "AI-контент, TG-бот",
    "Бухгалтерские бюро":           "AI-контент, TG-бот",
    "IT-компании / разработка":     "AI-контент, TG-бот (команда)",
    "Рекрутинговые агентства":      "AI-контент, TG-бот, автоматизация лидов",
    # Коучинг
    "Психологи (частные)":          "AI-контент, TG-бот, автоворонка",
    "Психотерапевты":               "AI-контент, TG-бот, автоворонка",
    "Коучи":                        "AI-контент, автопостинг, TG-бот, автоворонка",
    "Нутрициологи":                 "AI-контент, автопостинг, TG-бот, автоворонка",
    "Астрологи / таро":             "AI-контент, TG-бот",
    # Event и творчество
    "Флористы":                     "AI-контент, автопостинг, TG-бот",
    "Квесты":                       "AI-контент, TG-бот, автоматизация бронирования",
    "Кондитеры (торты на заказ)":  "AI-контент, автопостинг, TG-бот",
    "Организаторы свадеб":          "AI-контент, TG-бот, CRM-интеграция",
    "Фотографы (коммерческие)":     "AI-контент, TG-бот",
    "Event-агентства":              "AI-контент, TG-бот, CRM-интеграция",
    "Частные музеи и галереи":      "AI-контент, TG-бот",
    # Туризм
    "Туристические агентства":      "AI-контент, TG-бот, CRM-интеграция",
    "Прокат снаряжения":            "AI-контент, TG-бот",
    "Частные гиды":                 "AI-контент, TG-бот",
    # Зоо
    "Зоогостиницы (передержка)":    "AI-контент, TG-бот, автоматизация бронирования",
}

# ── Загрузка данных ─────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(FILE)
ws = wb.active

# Читаем заголовки строк 1-2
title_row1 = {}
title_row2 = {}
for cell in ws[1]:
    if cell.value:
        title_row1[cell.column] = {
            "value": cell.value, "font": cell.font, "fill": cell.fill,
            "alignment": cell.alignment, "border": cell.border
        }
for cell in ws[2]:
    if cell.value:
        title_row2[cell.column] = {
            "value": cell.value, "font": cell.font, "fill": cell.fill,
            "alignment": cell.alignment, "border": cell.border
        }

# Читаем все данные ниш (строки 3..last, только там где есть Ниша и Приоритет P0/P1/P2)
PRIO_ORDER = {"P0": 0, "P1": 1, "P2": 2}
data_rows = []

for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
    niche = row[0].value
    prio  = row[10].value if len(row) > 10 else None
    score = row[9].value  if len(row) > 9  else None
    if not niche or prio not in PRIO_ORDER:
        continue
    data_rows.append({
        "niche":    niche,
        "cat":      row[1].value,
        "scenario": row[2].value,
        "ai":       row[3].value,
        "auto":     row[4].value,
        "pay":      row[5].value,
        "budget":   row[6].value,
        "freq":     row[7].value,
        "comp":     row[8].value,
        "score":    score,
        "prio":     prio,
        "comment":  row[11].value,
        "integrations": row[12].value if len(row) > 12 else None,
    })

# Сортировка: P0→P1→P2, внутри по score убывание
data_rows.sort(key=lambda r: (PRIO_ORDER.get(r["prio"], 9), -(r["score"] or 0)))

# ── Сохраняем параметры столбцов ─────────────────────────────────────────────
col_widths = {i: ws.column_dimensions[get_column_letter(i)].width for i in range(1, ws.max_column + 1)}

# Легенда в конце
legend_data = []
for row in ws.iter_rows(min_row=ws.max_row - 10, max_row=ws.max_row):
    v0 = row[0].value
    if v0 and ("Приоритет" in str(v0) or str(v0).startswith("P")):
        legend_data.append([cell.value for cell in row])

# ── Пересобираем лист ────────────────────────────────────────────────────────
ws.delete_rows(1, ws.max_row)

# Заголовок строка 1 (merge)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)
h1 = ws.cell(row=1, column=1,
    value="and36.studio — Оценка ниш: платёжеспособность, сценарий услуг и приоритет")
h1.font = Font(bold=True, size=13, color="FFFFFF")
h1.fill = fill("1A1A2E")
h1.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 28

# Заголовки строка 2
HEADERS = [
    (1,  "Ниша",                        "left"),
    (2,  "Категория",                   "center"),
    (3,  "Сценарий\nуслуг",             "center"),
    (4,  "AI\nконтент",                 "center"),
    (5,  "Авто-\nматизация",            "center"),
    (6,  "Платёжеспо-\nсобность",       "center"),
    (7,  "Бюджет\n₽/мес",              "center"),
    (8,  "Частота\nконтакта",           "center"),
    (9,  "Конкуренция\nагентств",       "center"),
    (10, "Итог\n/10",                   "center"),
    (11, "Прио-\nритет",                "center"),
    (12, "Комментарий",                 "left"),
    (13, "Интеграции\n(сервисы)",       "left"),
    (14, "Предлагаемые\nуслуги",        "left"),
]
for col, title, align in HEADERS:
    cell = ws.cell(row=2, column=col, value=title)
    cell.font = Font(bold=True, size=10, color="FFFFFF")
    cell.fill = fill("16213E")
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 36

# ── Данные ───────────────────────────────────────────────────────────────────
for i, d in enumerate(data_rows):
    row_num = i + 3
    row_bg  = "F8F8F8" if i % 2 == 0 else "FFFFFF"
    cat     = d["cat"] or ""
    cat_col = CAT_COLORS.get(cat, "607D8B")
    service = SERVICES.get(d["niche"], "AI-контент, TG-бот")

    cells = [
        (1,  d["niche"],       row_bg,               True,  "left",   False),
        (2,  cat,              cat_col,               True,  "center", False),
        (3,  d["scenario"],    scenario_color(d["scenario"]), True, "center", False),
        (4,  d["ai"],          stars_color(d["ai"]),  False, "center", False),
        (5,  d["auto"],        stars_color(d["auto"]), False, "center", False),
        (6,  d["pay"],         pay_color(d["pay"]),   True,  "center", False),
        (7,  d["budget"],      row_bg,                False, "center", False),
        (8,  d["freq"],        freq_color(d["freq"]), False, "center", False),
        (9,  d["comp"],        comp_color(d["comp"]), False, "center", False),
        (10, d["score"],       score_color(d["score"]), True, "center", False),
        (11, d["prio"],        prio_color(d["prio"]), True,  "center", False),
        (12, d["comment"],     row_bg,                False, "left",   True),
        (13, d["integrations"],row_bg,                False, "left",   True),
        (14, service,          row_bg,                False, "left",   True),
    ]
    for col, value, bg, bold, align, wrap in cells:
        c = ws.cell(row=row_num, column=col, value=value)
        c.fill = fill(bg)
        c.font = Font(bold=bold, size=10, color="FFFFFF" if bg not in ("F8F8F8","FFFFFF","E8EAF6","E3F2FD","E8F5E9","FFF3E0","FFEBEE","FFF8E1","FFFDE7","E8F5E9","FFF8E1","E8F5E9","FFEBEE","2E7D32") else "1A1A2E")
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    ws.row_dimensions[row_num].height = 38

# ── Легенда ──────────────────────────────────────────────────────────────────
legend_start = len(data_rows) + 4
ws.cell(row=legend_start, column=1, value="")
leg_row = legend_start + 1
leg_items = [
    ("Приоритет", "Итоговый балл", "Описание", "161616", "FFFFFF"),
    ("P0 — Основные ниши",     "9–10", "Отличный fit, запускаем первыми", "1B5E20", "FFFFFF"),
    ("P1 — Расширение",        "7–8",  "Хороший fit, второй эшелон",      "E65100", "FFFFFF"),
    ("P2 — Будущее / нишевый", "5–6",  "Средний fit, при наличии ресурса","37474F", "FFFFFF"),
]
for label, score_val, desc, bg_c, fg_c in leg_items:
    c1 = ws.cell(row=leg_row, column=1, value=label)
    c1.font = Font(bold=True, size=10, color=fg_c)
    c1.fill = fill(bg_c)
    c1.alignment = Alignment(horizontal="left", vertical="center")
    c2 = ws.cell(row=leg_row, column=2, value=score_val)
    c2.font = Font(bold=True, size=10, color=fg_c)
    c2.fill = fill(bg_c)
    c2.alignment = Alignment(horizontal="center", vertical="center")
    c3 = ws.cell(row=leg_row, column=3, value=desc)
    c3.font = Font(size=10, color=fg_c)
    c3.fill = fill(bg_c)
    c3.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[leg_row].height = 22
    leg_row += 1

# ── Ширины столбцов ──────────────────────────────────────────────────────────
COL_WIDTHS = {1:32, 2:14, 3:10, 4:10, 5:10, 6:12, 7:10, 8:11, 9:13, 10:8, 11:8, 12:38, 13:38, 14:38}
for col, w in COL_WIDTHS.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# ── AutoFilter на строку 2 ───────────────────────────────────────────────────
last_data_row = len(data_rows) + 2
ws.auto_filter.ref = f"A2:{get_column_letter(14)}{last_data_row}"

# ── Закрепить строки 1–2 ─────────────────────────────────────────────────────
ws.freeze_panes = "A3"

wb.save(FILE)
print(f"✓ Таблица пересобрана: {len(data_rows)} ниш, 14 столбцов")
print(f"  Сортировка: P0={sum(1 for r in data_rows if r['prio']=='P0')} | "
      f"P1={sum(1 for r in data_rows if r['prio']=='P1')} | "
      f"P2={sum(1 for r in data_rows if r['prio']=='P2')}")
print(f"  AutoFilter: A2:{get_column_letter(14)}{last_data_row}")
print(f"  Freeze: строки 1–2")
