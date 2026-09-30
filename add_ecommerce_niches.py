import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

FILE = "niches_analysis.xlsx"

# New niches for "Парсинг товаров + автоконтент" service
# Format: (Ниша, Категория, Сценарий, AI, Авто, Платёжеспособность, Бюджет, Частота, Конкуренция, Итог, Приоритет, Комментарий, Интеграции)
NEW_NICHES = [
    (
        "Маркетплейс-продавцы (WB / Ozon)",
        "E-commerce",
        "А+Б",
        "★★★★★", "★★★★★",
        "Средняя", "20–45К", "Высокая", "Средняя",
        9, "P0",
        "Парсинг каталога WB/Ozon → посты в соцсети автоматически. Миллионы продавцов в РФ",
        "Wildberries API, Ozon API, amoCRM, WA/TG-бот"
    ),
    (
        "Инт. магазины (мебель и декор)",
        "E-commerce",
        "А+Б",
        "★★★★★", "★★★★★",
        "Средняя", "30–70К", "Средняя", "Средняя",
        8, "P0",
        "Визуальный продукт. Парсинг коллекций → интерьерные образы и истории для соцсетей",
        "1С / МойСклад, amoCRM, WA/TG-бот"
    ),
    (
        "Инт. магазины (ювелирные украшения)",
        "E-commerce",
        "А+Б",
        "★★★★★", "★★★★★",
        "Высокая", "35–80К", "Средняя", "Средняя",
        8, "P0",
        "Каждое изделие — история. Парсинг каталога → пост с эмоцией и описанием камня/металла",
        "1С / МойСклад, amoCRM, WA/TG-бот"
    ),
    (
        "Инт. магазины (электроника и гаджеты)",
        "E-commerce",
        "А+Б",
        "★★★★★", "★★★★★",
        "Высокая", "40–80К", "Средняя", "Высокая",
        7, "P1",
        "Частые новинки, технические характеристики. Парсинг → продающий контент без копирайтера",
        "1С / МойСклад, amoCRM, WA/TG-бот"
    ),
    (
        "Инт. магазины (товары для дома)",
        "E-commerce",
        "А+Б",
        "★★★★★", "★★★★★",
        "Средняя", "25–50К", "Средняя", "Высокая",
        7, "P1",
        "Большой каталог. Парсинг товаров → лайфстайл-контент и подборки для соцсетей",
        "1С / МойСклад, amoCRM, WA/TG-бот"
    ),
    (
        "Фермерские магазины / натуральная еда",
        "E-commerce",
        "А+Б",
        "★★★★★", "★★★★☆",
        "Средняя", "15–30К", "Высокая", "Низкая",
        7, "P1",
        "Сезонный каталог, эмоциональный продукт. Парсинг сезона → stories и посты с историей",
        "amoCRM, WA/TG-бот, Яндекс Еда"
    ),
    (
        "Аптеки и товары для здоровья онлайн",
        "E-commerce",
        "А+Б",
        "★★★★☆", "★★★★★",
        "Высокая", "40–100К", "Средняя", "Средняя",
        7, "P1",
        "Огромные каталоги. Парсинг описаний → образовательный контент без медицинских нарушений",
        "1С / МойСклад, amoCRM, WA/TG-бот"
    ),
    (
        "Книжные магазины онлайн",
        "E-commerce",
        "Б",
        "★★★★★", "★★★★★",
        "Низкая", "15–25К", "Высокая", "Низкая",
        6, "P2",
        "Каждая новинка — повод для поста. Парсинг аннотаций → обзоры, рекомендации, подборки",
        "1С / МойСклад, WA/TG-бот"
    ),
]

wb = openpyxl.load_workbook(FILE)
ws = wb.active

# Find the empty separator row before legend (row 107)
insert_before = None
for row in ws.iter_rows(min_row=100, max_row=ws.max_row):
    if row[0].value is None and row[1].value is None:
        insert_before = row[0].row
        break

if insert_before is None:
    insert_before = ws.max_row

# Get reference row styles from existing e-commerce rows (rows 72-74)
ref_rows = {}
for r in [72, 73, 74]:
    ref_rows[r] = [ws.cell(row=r, column=c) for c in range(1, 14)]

# Color palette matching the table
FILLS = {
    "row_odd":        "F8F8F8",
    "row_even":       "FFFFFF",
    "cat_ecommerce":  "E65100",  # orange category
    "scenario_ab":    "E8EAF6",
    "scenario_b":     "E3F2FD",
    "stars_5":        "1B5E20",
    "stars_4":        "2E7D32",
    "stars_3":        "E65100",
    "pay_high":       "E8F5E9",
    "pay_mid":        "FFF3E0",
    "pay_low":        "FFEBEE",
    "freq_high":      "E8F5E9",
    "freq_mid":       "FFF8E1",
    "freq_low":       "FFFDE7",
    "comp_high":      "FFEBEE",
    "comp_mid":       "FFF8E1",
    "comp_low":       "E8F5E9",
    "score_9":        "1B5E20",
    "score_8":        "2E7D32",
    "score_7":        "2E7D32",
    "score_6":        "E65100",
    "prio_p0":        "1B5E20",
    "prio_p1":        "E65100",
    "prio_p2":        "37474F",
}

def make_fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)

def stars_fill(stars_str):
    count = stars_str.count("★")
    if count == 5: return FILLS["stars_5"]
    if count == 4: return FILLS["stars_4"]
    return FILLS["stars_3"]

def pay_fill(val):
    if val == "Высокая": return FILLS["pay_high"]
    if val == "Средняя": return FILLS["pay_mid"]
    return FILLS["pay_low"]

def freq_fill(val):
    if val == "Высокая": return FILLS["freq_high"]
    if val == "Средняя": return FILLS["freq_mid"]
    return FILLS["freq_low"]

def comp_fill(val):
    if val == "Высокая": return FILLS["comp_high"]
    if val == "Средняя": return FILLS["comp_mid"]
    return FILLS["comp_low"]

def score_fill(val):
    if val >= 9: return FILLS["score_9"]
    if val >= 8: return FILLS["score_8"]
    if val >= 7: return FILLS["score_7"]
    return FILLS["score_6"]

def prio_fill(val):
    if val == "P0": return FILLS["prio_p0"]
    if val == "P1": return FILLS["prio_p1"]
    return FILLS["prio_p2"]

# Insert rows
ws.insert_rows(insert_before, amount=len(NEW_NICHES))

for i, niche in enumerate(NEW_NICHES):
    row_num = insert_before + i
    row_bg = "F8F8F8" if i % 2 == 0 else "FFFFFF"

    (name, cat, scenario, ai, auto, pay, budget, freq, comp, score, prio, comment, integrations) = niche

    cells_data = [
        (1, name,         row_bg,              True,  "left"),
        (2, cat,          FILLS["cat_ecommerce"], True, "center"),
        (3, scenario,     FILLS["scenario_ab"] if "А" in scenario else FILLS["scenario_b"], True, "center"),
        (4, ai,           stars_fill(ai),      False, "center"),
        (5, auto,         stars_fill(auto),    False, "center"),
        (6, pay,          pay_fill(pay),        True,  "center"),
        (7, budget,       row_bg,              False, "center"),
        (8, freq,         freq_fill(freq),     False, "center"),
        (9, comp,         comp_fill(comp),     False, "center"),
        (10, score,       score_fill(score),   True,  "center"),
        (11, prio,        prio_fill(prio),     True,  "center"),
        (12, comment,     row_bg,              False, "left"),
        (13, integrations,"000000",            False, "left"),
    ]

    for col, value, fill_hex, bold, align in cells_data:
        cell = ws.cell(row=row_num, column=col, value=value)
        if fill_hex != "000000":
            cell.fill = make_fill(fill_hex)
        else:
            cell.fill = PatternFill(fill_type=None)
        cell.font = Font(bold=bold, size=10)
        cell.alignment = Alignment(
            horizontal=align,
            vertical="center",
            wrap_text=(col in [12, 13])
        )

# Set row height for new rows
for i in range(len(NEW_NICHES)):
    ws.row_dimensions[insert_before + i].height = 35

wb.save(FILE)
print(f"✓ Добавлено {len(NEW_NICHES)} новых ниш E-commerce в {FILE}")
for n in NEW_NICHES:
    print(f"  — {n[0]} ({n[10]})")
